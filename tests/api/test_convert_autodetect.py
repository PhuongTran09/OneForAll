import io

import docx
import pytest
from httpx import AsyncClient
from PIL import Image

from app.repositories.job_repository import JobRepository
from app.worker.tasks.convert import process_convert_job
from tests.conftest import create_test_supabase_token


def _auth_headers() -> dict[str, str]:
    token = create_test_supabase_token(
        email="converter-test@example.com",
        username="converteruser",
        full_name="Converter User",
    )
    return {"Authorization": f"Bearer {token}"}


def _create_sample_png() -> bytes:
    img = Image.new("RGB", (32, 32), color="blue")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _create_sample_jpg() -> bytes:
    img = Image.new("RGB", (32, 32), color="green")
    buf = io.BytesIO()
    img.save(buf, format="JPEG")
    return buf.getvalue()


def _create_sample_docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("Autodetect Test Document", level=1)
    doc.add_paragraph("This is an autodetected docx document for testing.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_convert_autodetect_txt_to_pdf(client: AsyncClient):
    """Test 1: txt -> pdf without from_format, end-to-end with worker."""
    headers = _auth_headers()
    txt_content = b"OneForAll AutoDetect Test\nLine 2 of document content."

    response = await client.post(
        "/api/v1/convert-file",
        headers=headers,
        data={"to": "pdf"},
        files={"file": ("notes.txt", txt_content, "text/plain")},
    )
    assert response.status_code == 202
    job_id = response.json()["job_id"]

    repo = JobRepository()
    job = await repo.get(job_id)
    assert job is not None
    assert job.job_metadata["from"] == "txt"
    assert job.job_metadata["to"] == "pdf"

    # Worker processing
    result = process_convert_job(job_id)
    assert result["status"] == "completed"
    assert result["output_key"] == f"outputs/{job_id}/result.pdf"


@pytest.mark.asyncio
async def test_convert_autodetect_docx_to_pdf(client: AsyncClient):
    """Test 2: docx -> pdf without from_format, end-to-end with worker."""
    headers = _auth_headers()
    docx_bytes = _create_sample_docx()

    response = await client.post(
        "/api/v1/convert-file",
        headers=headers,
        data={"to": "pdf"},
        files={
            "file": (
                "document.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 202
    job_id = response.json()["job_id"]

    repo = JobRepository()
    job = await repo.get(job_id)
    assert job is not None
    assert job.job_metadata["from"] == "docx"
    assert job.job_metadata["to"] == "pdf"

    # Worker processing
    result = process_convert_job(job_id)
    assert result["status"] == "completed"
    assert result["output_key"] == f"outputs/{job_id}/result.pdf"


@pytest.mark.asyncio
async def test_convert_autodetect_png_to_jpg(client: AsyncClient):
    """Test 3: png -> jpg without from_format, end-to-end with worker."""
    headers = _auth_headers()
    png_bytes = _create_sample_png()

    response = await client.post(
        "/api/v1/convert-file",
        headers=headers,
        data={"to": "jpg"},
        files={"file": ("picture.png", png_bytes, "image/png")},
    )
    assert response.status_code == 202
    job_id = response.json()["job_id"]

    repo = JobRepository()
    job = await repo.get(job_id)
    assert job is not None
    assert job.job_metadata["from"] == "png"
    assert job.job_metadata["to"] == "jpg"

    # Worker processing
    result = process_convert_job(job_id)
    assert result["status"] == "completed"
    assert result["output_key"] == f"outputs/{job_id}/result.jpg"


@pytest.mark.asyncio
async def test_convert_autodetect_jpg_to_png(client: AsyncClient):
    """Test 4: jpg -> png without from_format, end-to-end with worker."""
    headers = _auth_headers()
    jpg_bytes = _create_sample_jpg()

    response = await client.post(
        "/api/v1/convert-file",
        headers=headers,
        data={"to": "png"},
        files={"file": ("photo.jpg", jpg_bytes, "image/jpeg")},
    )
    assert response.status_code == 202
    job_id = response.json()["job_id"]

    repo = JobRepository()
    job = await repo.get(job_id)
    assert job is not None
    assert job.job_metadata["from"] == "jpg"
    assert job.job_metadata["to"] == "png"

    # Worker processing
    result = process_convert_job(job_id)
    assert result["status"] == "completed"
    assert result["output_key"] == f"outputs/{job_id}/result.png"


@pytest.mark.asyncio
async def test_convert_unsupported_formats(client: AsyncClient):
    """Test 5: trường hợp format không hỗ trợ."""
    headers = _auth_headers()

    # 5a. File format cannot be determined (random binary with no extension and generic MIME)
    unknown_bytes = b"\x00\x01\x02\x03\x04\x05\x06\x07\xfe\xff"
    res_unknown = await client.post(
        "/api/v1/convert-file",
        headers=headers,
        data={"to": "pdf"},
        files={"file": ("unknown_file", unknown_bytes, "application/octet-stream")},
    )
    assert res_unknown.status_code == 400
    assert "Không thể xác định định dạng" in res_unknown.json()["detail"]

    # 5b. Valid detected format (png), but conversion pair (png -> docx) is unsupported
    png_bytes = _create_sample_png()
    res_pair = await client.post(
        "/api/v1/convert-file",
        headers=headers,
        data={"to": "docx"},
        files={"file": ("image.png", png_bytes, "image/png")},
    )
    assert res_pair.status_code == 400
    assert "chưa được hỗ trợ" in res_pair.json()["detail"]


@pytest.mark.asyncio
async def test_convert_no_extension_magic_bytes_detection(client: AsyncClient):
    """Test 6: File has no extension and generic MIME, detected via magic bytes."""
    headers = _auth_headers()
    png_bytes = _create_sample_png()

    response = await client.post(
        "/api/v1/convert-file",
        headers=headers,
        data={"to": "jpg"},
        files={"file": ("raw_upload_blob", png_bytes, "application/octet-stream")},
    )
    assert response.status_code == 202
    job_id = response.json()["job_id"]

    repo = JobRepository()
    job = await repo.get(job_id)
    assert job is not None
    assert job.job_metadata["from"] == "png"
    assert job.job_metadata["to"] == "jpg"

    result = process_convert_job(job_id)
    assert result["status"] == "completed"


@pytest.mark.asyncio
async def test_convert_normalizes_to_format(client: AsyncClient):
    """Test 7: to_format with leading dot and uppercase (.JPG) is normalized."""
    headers = _auth_headers()
    png_bytes = _create_sample_png()

    response = await client.post(
        "/api/v1/convert-file",
        headers=headers,
        data={"to": ".JPG"},
        files={"file": ("sample.png", png_bytes, "image/png")},
    )
    assert response.status_code == 202
    job_id = response.json()["job_id"]

    repo = JobRepository()
    job = await repo.get(job_id)
    assert job is not None
    assert job.job_metadata["to"] == "jpg"
