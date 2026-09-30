import io

import docx
import openpyxl
import pytest
from httpx import AsyncClient
from PIL import Image


def _create_sample_png() -> bytes:
    img = Image.new("RGBA", (50, 50), (255, 0, 0, 128))
    bio = io.BytesIO()
    img.save(bio, format="PNG")
    return bio.getvalue()


def _create_sample_jpg() -> bytes:
    img = Image.new("RGB", (50, 50), (0, 128, 255))
    bio = io.BytesIO()
    img.save(bio, format="JPEG")
    return bio.getvalue()


def _create_sample_docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("Tiêu đề thử nghiệm", level=1)
    doc.add_paragraph("Nội dung tài liệu tiếng Việt có dấu: Kiểm tra chuyển đổi sang PDF.")
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def _create_sample_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Employees"
    ws.append(["id", "name", "role"])
    ws.append([101, "Nguyễn Văn A", "Developer"])
    ws.append([102, "Trần Thị B", "Designer"])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


@pytest.mark.asyncio
async def test_png_to_svg(client: AsyncClient):
    png_bytes = _create_sample_png()
    response = await client.post(
        "/api/v1/convert/png-to-svg",
        files={"file": ("test.png", png_bytes, "image/png")},
        data={"colormode": "color", "mode": "spline"},
    )
    assert response.status_code == 200
    assert "image/svg+xml" in response.headers["content-type"]
    assert b"<svg" in response.content or b"<?xml" in response.content


@pytest.mark.asyncio
async def test_doc_to_pdf(client: AsyncClient):
    docx_bytes = _create_sample_docx()
    response = await client.post(
        "/api/v1/convert/doc-to-pdf",
        files={
            "file": (
                "sample.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 200
    assert "application/pdf" in response.headers["content-type"]
    assert response.content.startswith(b"%PDF-")


@pytest.mark.asyncio
async def test_txt_to_pdf(client: AsyncClient):
    txt_content = "Báo cáo tiến độ dự án OneForAll.\nTiếng Việt: Hà Nội, Đà Nẵng, TP. Hồ Chí Minh."
    response = await client.post(
        "/api/v1/convert/txt-to-pdf",
        files={"file": ("report.txt", txt_content.encode("utf-8"), "text/plain")},
        data={"title": "Báo cáo OneForAll"},
    )
    assert response.status_code == 200
    assert "application/pdf" in response.headers["content-type"]
    assert response.content.startswith(b"%PDF-")


@pytest.mark.asyncio
async def test_txt_to_doc(client: AsyncClient):
    txt_content = "Đoạn văn bản 1\nĐoạn văn bản 2"
    response = await client.post(
        "/api/v1/convert/txt-to-doc",
        files={"file": ("notes.txt", txt_content.encode("utf-8"), "text/plain")},
        data={"title": "Ghi chú"},
    )
    assert response.status_code == 200
    assert (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        in response.headers["content-type"]
    )
    assert response.content.startswith(b"PK")  # Zip / DOCX header


@pytest.mark.asyncio
async def test_xlsx_to_csv(client: AsyncClient):
    xlsx_bytes = _create_sample_xlsx()
    response = await client.post(
        "/api/v1/convert/xlsx-to-csv",
        files={
            "file": (
                "data.xlsx",
                xlsx_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200
    assert "text/csv" in response.headers["content-type"]
    csv_text = response.content.decode("utf-8-sig")
    assert "Nguyễn Văn A" in csv_text
    assert "Designer" in csv_text


@pytest.mark.asyncio
async def test_xlsx_to_json(client: AsyncClient):
    xlsx_bytes = _create_sample_xlsx()
    response = await client.post(
        "/api/v1/convert/xlsx-to-json",
        files={
            "file": (
                "data.xlsx",
                xlsx_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200
    data = response.json()
    assert isinstance(data, list)
    assert len(data) == 2
    assert data[0]["name"] == "Nguyễn Văn A"
    assert data[1]["role"] == "Designer"


@pytest.mark.asyncio
async def test_png_to_jpg(client: AsyncClient):
    png_bytes = _create_sample_png()
    response = await client.post(
        "/api/v1/convert/png-to-jpg",
        files={"file": ("transparent.png", png_bytes, "image/png")},
        data={"quality": 90, "bg_color": "#ffffff"},
    )
    assert response.status_code == 200
    assert "image/jpeg" in response.headers["content-type"]
    assert response.content.startswith(b"\xff\xd8")  # JPEG header


@pytest.mark.asyncio
async def test_jpg_to_webp(client: AsyncClient):
    jpg_bytes = _create_sample_jpg()
    response = await client.post(
        "/api/v1/convert/jpg-to-webp",
        files={"file": ("photo.jpg", jpg_bytes, "image/jpeg")},
        data={"quality": 85},
    )
    assert response.status_code == 200
    assert "image/webp" in response.headers["content-type"]
    assert response.content.startswith(b"RIFF")
    assert b"WEBP" in response.content[:16]


@pytest.mark.asyncio
async def test_empty_file_error(client: AsyncClient):
    response = await client.post(
        "/api/v1/convert/png-to-svg",
        files={"file": ("empty.png", b"", "image/png")},
    )
    assert response.status_code == 400


@pytest.mark.asyncio
async def test_txt_to_pdf_with_vietnamese_filename(client: AsyncClient):
    txt_content = "Nội dung tài liệu với tên file tiếng Việt."
    response = await client.post(
        "/api/v1/convert/txt-to-pdf",
        files={"file": ("tài_liệu_phụ.txt", txt_content.encode("utf-8"), "text/plain")},
        data={"title": "Tài liệu phụ"},
    )
    assert response.status_code == 200
    assert "application/pdf" in response.headers["content-type"]
    assert "content-disposition" in response.headers
    # Check ASCII fallback and RFC 5987 encoded name in header
    disposition = response.headers["content-disposition"]
    assert 'filename="tai_lieu_phu.pdf"' in disposition
    assert "filename*=UTF-8''" in disposition

