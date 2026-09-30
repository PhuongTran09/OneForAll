"""File & Document Conversion Service.

Provides format conversion for:
- PNG to SVG (vector tracing via vtracer)
- DOCX to PDF (semantic HTML extraction + PDF rendering)
- TXT to PDF & TXT to DOCX
- XLSX to CSV & XLSX to JSON
- PNG to JPG/JPEG
- JPG to WEBP

Lưu ý: các hàm ở đây là hàm chặn (blocking, CPU-bound). Trong FastAPI hãy gọi qua
`await run_in_threadpool(file_converter_service.xxx, ...)` hoặc dùng endpoint `def` thường.
"""

import codecs
import csv
import functools
import html
import io
import os
import re
import unicodedata
from collections.abc import Callable
from datetime import date, time, timedelta
from pathlib import Path
from string import Template
from typing import Any, ParamSpec, TypeVar
from urllib.parse import quote

from PIL import Image, ImageOps

from app.core.config import settings
from app.core.exceptions import AppException
from app.utils.logger import logger

P = ParamSpec("P")
R = TypeVar("R")

# Giới hạn số pixel của ảnh đầu vào (tránh decompression bomb / hết RAM)
MAX_IMAGE_PIXELS: int = getattr(settings, "MAX_IMAGE_PIXELS", 50_000_000)

FONT_CANDIDATES: tuple[str, ...] = (
    "C:/Windows/Fonts/arial.ttf",
    "C:/Windows/Fonts/segoeui.ttf",
    "C:/Windows/Fonts/tahoma.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    "/usr/share/fonts/TTF/DejaVuSans.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/Library/Fonts/Arial.ttf",
)

VTRACER_COLORMODES = frozenset({"color", "binary"})
VTRACER_MODES = frozenset({"spline", "polygon", "none"})

# Ký tự điều khiển không hợp lệ trong XML (python-docx sẽ ném ValueError) và vô nghĩa trong PDF
_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")

# CSS chỉ được dựng 1 lần thay vì mỗi lần gọi
_CSS_TEMPLATE = Template(
    """
@page {
    size: a4 portrait;
    margin: 2cm 1.5cm 2cm 1.5cm;
    @bottom-right {
        content: counter(page);
        font-size: 9pt;
        color: #6b7280;
    }
}
$font_face
* {
    font-family: $font_family;
    box-sizing: border-box;
}
body {
    font-size: 11pt;
    line-height: 1.6;
    color: #1f2937;
    margin: 0;
    padding: 0;
}
h1, h2, h3, h4, h5, h6 {
    color: #111827;
    margin-top: 1.2em;
    margin-bottom: 0.5em;
}
h1 { font-size: 18pt; }
h2 { font-size: 15pt; }
h3 { font-size: 13pt; }
p {
    margin: 0 0 0.8em 0;
    text-align: justify;
}
table {
    border-collapse: collapse;
    width: 100%;
    margin: 1em 0;
    font-size: 10pt;
}
th, td {
    border: 1px solid #d1d5db;
    padding: 6px 10px;
    text-align: left;
    vertical-align: top;
}
th {
    background-color: #f3f4f6;
    font-weight: bold;
    color: #111827;
}
tr:nth-child(even) td {
    background-color: #f9fafb;
}
pre {
    background-color: #f8fafc;
    border: 1px solid #e2e8f0;
    border-radius: 4px;
    padding: 10px;
    font-family: $font_family;
    font-size: 9.5pt;
    line-height: 1.5;
    white-space: pre-wrap;
    word-break: break-word;
}
ul, ol {
    margin: 0 0 1em 1.5em;
    padding: 0;
}
li {
    margin-bottom: 0.3em;
}
img {
    max-width: 100%;
    height: auto;
}
"""
)


def _handle_errors(action: str) -> Callable[[Callable[P, R]], Callable[P, R]]:
    """Gom logic bắt lỗi lặp lại ở mọi hàm chuyển đổi.

    - AppException (đã có status code chuẩn) được giữ nguyên, không bị bọc lại thành 400.
    - Thiếu thư viện -> 500 (lỗi server, không phải lỗi của người dùng).
    - Lỗi còn lại (file hỏng, sai định dạng...) -> 400.
    """

    def decorator(func: Callable[P, R]) -> Callable[P, R]:
        @functools.wraps(func)
        def wrapper(*args: P.args, **kwargs: P.kwargs) -> R:
            try:
                return func(*args, **kwargs)
            except AppException:
                raise
            except ImportError as e:
                logger.exception(f"[!] Thiếu thư viện khi {action}")
                raise AppException(
                    status_code=500, detail=f"Máy chủ thiếu thư viện cần thiết để {action}."
                ) from e
            except Exception as e:
                logger.exception(f"[!] Lỗi {action}")
                raise AppException(status_code=400, detail=f"Lỗi {action}: {e!s}") from e

        return wrapper

    return decorator


def _decode_text(data: str | bytes) -> str:
    """Giải mã văn bản (tự đoán encoding) và làm sạch ký tự không hợp lệ."""
    if isinstance(data, bytes):
        if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
            text = data.decode("utf-16", errors="replace")
        else:
            # utf-8-sig xử lý được cả UTF-8 thường lẫn có BOM (utf-8 thuần sẽ để lại ký tự \ufeff)
            for enc in ("utf-8-sig", "cp1258"):
                try:
                    text = data.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
            else:
                text = data.decode("latin-1")  # không bao giờ lỗi
    else:
        text = data

    text = text.replace("\r\n", "\n").replace("\r", "\n")
    return _CONTROL_CHARS.sub("", text)


class FileConverterService:
    """Service chịu trách nhiệm chuyển đổi đa định dạng hình ảnh và tài liệu."""

    def __init__(self):
        self._font_path: str | None = self._detect_system_font()
        self._css: str = self._build_css()
        self._pdf_policy: Any = None  # tạo lười (lazy) vì xhtml2pdf import khá nặng

    # ------------------------------------------------------------------ helpers
    @staticmethod
    def build_content_disposition(filename: str) -> str:
        normalized = unicodedata.normalize("NFKD", filename)

        ascii_filename = normalized.encode(
            "ascii",
            "ignore",
        ).decode("ascii")

        if not ascii_filename:
            ascii_filename = "download"

        encoded_filename = quote(
            filename,
            safe="",
        )

        return (
            f'attachment; filename="{ascii_filename}"; '
            f"filename*=UTF-8''{encoded_filename}"
        )
  
    @staticmethod
    def _detect_system_font() -> str | None:
        """Tìm font hệ thống có hỗ trợ Unicode/Tiếng Việt tốt nhất."""
        for candidate in FONT_CANDIDATES:
            if os.path.isfile(candidate):
                return candidate.replace("\\", "/")
        return None

    def _build_css(self) -> str:
        font_family = "Helvetica, Arial, sans-serif"
        font_face = ""
        if self._font_path:
            font_family = "'AppCustomFont', sans-serif"
            font_face = (
                "@font-face { font-family: 'AppCustomFont'; "
                f"src: url('{self._font_path}'); }}"
            )
        return _CSS_TEMPLATE.substitute(font_face=font_face, font_family=font_family)

    def _build_html_wrapper(self, body_content: str, title: str = "Tài liệu") -> str:
        """Bọc nội dung vào mẫu HTML chuẩn để xuất ra PDF chất lượng cao."""
        return (
            '<!DOCTYPE html>\n<html>\n<head>\n<meta charset="utf-8">\n'
            f"<title>{html.escape(title)}</title>\n<style>{self._css}</style>\n"
            f"</head>\n<body>\n{body_content}\n</body>\n</html>"
        )

    def _get_pdf_policy(self) -> Any:
        if self._pdf_policy is None:
            from xhtml2pdf.config.resources import ResourceAccessPolicy

            extra_roots = (Path(self._font_path).parent,) if self._font_path else ()
            self._pdf_policy = ResourceAccessPolicy(extra_roots=extra_roots)
        return self._pdf_policy

    def _render_html_to_pdf(self, html_content: str) -> bytes:
        """Biên dịch HTML sang PDF bằng xhtml2pdf với chính sách bảo mật tài nguyên."""
        from xhtml2pdf import pisa

        out = io.BytesIO()
        result = pisa.CreatePDF(
            html_content,
            dest=out,
            encoding="utf-8",
            resource_policy=self._get_pdf_policy(),
        )
        if result.err:
            logger.error(f"[!] Lỗi khi xuất file PDF: mã lỗi {result.err}")
            raise AppException(status_code=500, detail="Không thể tạo file PDF từ tài liệu.")
        return out.getvalue()

    @staticmethod
    def _open_image(data: bytes) -> Image.Image:
        """Mở ảnh (lazy) và chặn ảnh quá lớn TRƯỚC khi giải mã toàn bộ."""
        img = Image.open(io.BytesIO(data))
        if img.width * img.height > MAX_IMAGE_PIXELS:
            img.close()
            raise AppException(
                status_code=413,
                detail=f"Ảnh quá lớn (tối đa {MAX_IMAGE_PIXELS:,} pixel).",
            )
        return img

    @staticmethod
    def _load_workbook(xlsx_bytes: bytes):
        import openpyxl

        # read_only=True: đọc dạng stream, nhanh hơn và tốn RAM ít hơn nhiều với file lớn
        return openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)

    @staticmethod
    def _pick_sheet(wb, sheet_name: str | None):
        if not sheet_name:
            sheet = wb.active or wb.worksheets[0]
        elif sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
        else:
            raise AppException(status_code=400, detail=f"Không tìm thấy sheet '{sheet_name}'.")
        # Một số file xlsx khai báo sai vùng dữ liệu -> read_only chỉ đọc thiếu cột/dòng
        reset = getattr(sheet, "reset_dimensions", None)
        if reset:
            reset()
        return sheet

    @staticmethod
    def _plain_value(value: Any) -> Any:
        """Chuẩn hóa giá trị ô Excel thành kiểu serialize được (CSV/JSON)."""
        if isinstance(value, (date, time)):  # datetime là lớp con của date
            return value.isoformat()
        if isinstance(value, timedelta):
            return str(value)
        return value

    # -------------------------------------------------------------------------
    # 1. PNG sang SVG
    # -------------------------------------------------------------------------
    @_handle_errors("chuyển đổi PNG sang SVG")
    def png_to_svg(
        self,
        image_bytes: bytes,
        colormode: str = "color",
        mode: str = "spline",
    ) -> bytes:
        """Chuyển đổi hình ảnh PNG sang định dạng vector SVG bằng vtracer.

        :param image_bytes: Dữ liệu bytes của file PNG.
        :param colormode: 'color' hoặc 'binary' (đen trắng).
        :param mode: Kiểu đường nét 'spline', 'polygon', hoặc 'none'.
        :return: SVG bytes.
        """
        import vtracer

        if colormode not in VTRACER_COLORMODES:
            raise AppException(status_code=400, detail=f"colormode phải thuộc {sorted(VTRACER_COLORMODES)}.")
        if mode not in VTRACER_MODES:
            raise AppException(status_code=400, detail=f"mode phải thuộc {sorted(VTRACER_MODES)}.")

        with self._open_image(image_bytes) as img:
            if img.format != "PNG":
                # Chuẩn hóa về PNG nếu file gửi lên là định dạng khác
                if img.mode not in ("RGB", "RGBA", "L", "LA", "P"):
                    img = img.convert("RGBA")
                buf = io.BytesIO()
                img.save(buf, format="PNG")
                image_bytes = buf.getvalue()

        # Truyền thẳng bytes (không dùng list(image_bytes) - tạo list hàng triệu int rất chậm và tốn RAM)
        svg_str = vtracer.convert_raw_image_to_svg(
            image_bytes,
            img_format="png",
            colormode=colormode,
            mode=mode,
        )
        return svg_str.encode("utf-8")

    # -------------------------------------------------------------------------
    # 2. DOCX sang PDF
    # -------------------------------------------------------------------------
    @_handle_errors("khi chuyển đổi file DOC/DOCX sang PDF")
    def doc_to_pdf(self, doc_bytes: bytes) -> bytes:
        """Chuyển đổi tài liệu DOCX sang PDF.

        Chỉ hỗ trợ .docx (định dạng .doc cũ dùng OLE, mammoth không đọc được).

        :param doc_bytes: Dữ liệu bytes của file DOCX.
        :return: PDF bytes.
        """
        if not doc_bytes.startswith(b"PK"):  # .docx thực chất là file zip
            raise AppException(
                status_code=415,
                detail="Chỉ hỗ trợ .docx. Định dạng .doc (Word 97-2003) chưa được hỗ trợ, "
                "vui lòng lưu lại dưới dạng .docx.",
            )

        import mammoth

        result = mammoth.convert_to_html(io.BytesIO(doc_bytes))
        full_html = self._build_html_wrapper(result.value, title="Tài liệu DOCX")
        return self._render_html_to_pdf(full_html)

    # -------------------------------------------------------------------------
    # 3. TXT sang PDF
    # -------------------------------------------------------------------------
    @_handle_errors("chuyển đổi TXT sang PDF")
    def txt_to_pdf(self, text_content: str | bytes, title: str = "Tài liệu Text") -> bytes:
        """Chuyển đổi văn bản TXT sang file PDF.

        :param text_content: Chuỗi văn bản hoặc bytes của file TXT.
        :param title: Tiêu đề tài liệu.
        :return: PDF bytes.
        """
        text = _decode_text(text_content)
        body_html = f"<h1>{html.escape(title)}</h1>\n<pre>{html.escape(text)}</pre>"
        return self._render_html_to_pdf(self._build_html_wrapper(body_html, title=title))

    # -------------------------------------------------------------------------
    # 3b. TXT sang DOCX
    # -------------------------------------------------------------------------
    @_handle_errors("chuyển đổi TXT sang DOC")
    def txt_to_doc(self, text_content: str | bytes, title: str | None = None) -> bytes:
        """Chuyển đổi văn bản TXT sang định dạng DOCX.

        :param text_content: Chuỗi văn bản hoặc bytes của file TXT.
        :param title: Tiêu đề tài liệu (nếu có sẽ đặt Heading 1).
        :return: DOCX bytes.
        """
        import docx

        text = _decode_text(text_content)
        doc = docx.Document()
        if title:
            doc.add_heading(_CONTROL_CHARS.sub("", title), level=1)
        for line in text.split("\n"):
            doc.add_paragraph(line)

        out = io.BytesIO()
        doc.save(out)
        return out.getvalue()

    # -------------------------------------------------------------------------
    # 4a. XLSX sang CSV
    # -------------------------------------------------------------------------
    @_handle_errors("chuyển đổi XLSX sang CSV")
    def xlsx_to_csv(self, xlsx_bytes: bytes, sheet_name: str | None = None) -> bytes:
        """Chuyển đổi bảng tính XLSX sang CSV (UTF-8 with BOM để mở tốt trên Excel).

        :param xlsx_bytes: Dữ liệu bytes file XLSX.
        :param sheet_name: Tên sheet cần chuyển đổi (mặc định lấy active sheet).
        :return: CSV bytes.
        """
        wb = self._load_workbook(xlsx_bytes)
        try:
            sheet = self._pick_sheet(wb, sheet_name)
            output = io.StringIO()
            writer = csv.writer(output)  # tự chuyển None -> "" và mọi kiểu khác -> str

            pending_blank: list[list[Any]] = []  # dòng trống chờ ghi, bỏ các dòng trống ở cuối sheet
            for row in sheet.iter_rows(values_only=True):
                cells = [self._plain_value(c) for c in row]
                if all(c is None for c in cells):
                    pending_blank.append(cells)
                    continue
                if pending_blank:
                    writer.writerows(pending_blank)
                    pending_blank.clear()
                writer.writerow(cells)
        finally:
            wb.close()

        # utf-8-sig (có BOM) để Excel trên Windows hiển thị đúng tiếng Việt
        return output.getvalue().encode("utf-8-sig")

    # -------------------------------------------------------------------------
    # 4b. XLSX sang JSON
    # -------------------------------------------------------------------------
    @_handle_errors("chuyển đổi XLSX sang JSON")
    def xlsx_to_json(
        self,
        xlsx_bytes: bytes,
        sheet_name: str | None = None,
        all_sheets: bool = False,
    ) -> Any:
        """Chuyển đổi bảng tính XLSX sang định dạng JSON.

        :param xlsx_bytes: Dữ liệu bytes file XLSX.
        :param sheet_name: Tên sheet cần chuyển đổi (nếu all_sheets=False).
        :param all_sheets: Nếu True, trả về dict gồm toàn bộ các sheet.
        :return: Dict hoặc List chứa dữ liệu JSON.
        """
        wb = self._load_workbook(xlsx_bytes)
        try:
            if all_sheets:
                return {name: self._sheet_to_records(self._pick_sheet(wb, name)) for name in wb.sheetnames}
            return self._sheet_to_records(self._pick_sheet(wb, sheet_name))
        finally:
            wb.close()

    def _sheet_to_records(self, sheet) -> list[dict[str, Any]]:
        rows = sheet.iter_rows(values_only=True)
        first = next(rows, None)  # stream: không nạp cả sheet vào list
        if first is None:
            return []

        # Dòng đầu là header; ô trống -> col_N; header trùng tên được đánh số để không bị ghi đè
        headers: list[str] = []
        seen: dict[str, int] = {}
        for idx, h in enumerate(first):
            name = str(h) if h is not None else f"col_{idx}"
            if name in seen:
                seen[name] += 1
                name = f"{name}_{seen[name]}"
            else:
                seen[name] = 0
            headers.append(name)

        records: list[dict[str, Any]] = []
        for row in rows:
            if all(c is None for c in row):
                continue
            records.append(
                {
                    (headers[idx] if idx < len(headers) else f"col_{idx}"): self._plain_value(val)
                    for idx, val in enumerate(row)
                }
            )
        return records

    # -------------------------------------------------------------------------
    # 5. PNG sang JPG / JPEG
    # -------------------------------------------------------------------------
    @_handle_errors("chuyển đổi PNG sang JPG")
    def png_to_jpg(
        self,
        png_bytes: bytes,
        quality: int = 95,
        bg_color: tuple[int, int, int] = (255, 255, 255),
    ) -> bytes:
        """Chuyển đổi ảnh PNG sang JPG/JPEG, xử lý nền trong suốt bằng màu nền tùy chọn.

        :param png_bytes: Dữ liệu bytes ảnh PNG.
        :param quality: Chất lượng ảnh nén JPEG (1-100).
        :param bg_color: Màu RGB thay thế cho nền trong suốt (mặc định trắng).
        :return: JPEG bytes.
        """
        quality = max(1, min(100, quality))

        with self._open_image(png_bytes) as image:
            has_alpha = image.mode in ("RGBA", "LA", "PA") or (
                image.mode == "P" and "transparency" in image.info
            )
            if has_alpha:
                rgba = image.convert("RGBA")  # chỉ convert 1 lần
                rgb = Image.new("RGB", rgba.size, bg_color)
                rgb.paste(rgba, mask=rgba.getchannel("A"))
            else:
                rgb = image.convert("RGB")

            out = io.BytesIO()
            rgb.save(out, format="JPEG", quality=quality, optimize=True)
            return out.getvalue()

    # -------------------------------------------------------------------------
    # 6. JPG sang WEBP
    # -------------------------------------------------------------------------
    @_handle_errors("chuyển đổi JPG sang WEBP")
    def jpg_to_webp(
        self,
        jpg_bytes: bytes,
        quality: int = 90,
        lossless: bool = False,
        method: int = 4,
    ) -> bytes:
        """Chuyển đổi ảnh JPG/JPEG sang định dạng WEBP.

        :param jpg_bytes: Dữ liệu bytes ảnh JPG.
        :param quality: Mức chất lượng nén (1-100).
        :param lossless: Có nén không suy giảm chất lượng hay không.
        :param method: 0 (nhanh nhất) đến 6 (nén tốt nhất, chậm nhất). Mặc định 4.
        :return: WEBP bytes.
        """
        quality = max(1, min(100, quality))
        method = max(0, min(6, method))

        with self._open_image(jpg_bytes) as image:
            icc_profile = image.info.get("icc_profile")  # giữ hồ sơ màu, tránh lệch màu
            image = ImageOps.exif_transpose(image)  # ảnh điện thoại: xoay đúng chiều vì WEBP sẽ mất EXIF
            out = io.BytesIO()
            image.save(
                out,
                format="WEBP",
                quality=quality,
                lossless=lossless,
                method=method,
                icc_profile=icc_profile,
            )
            return out.getvalue()
    

file_converter_service = FileConverterService()